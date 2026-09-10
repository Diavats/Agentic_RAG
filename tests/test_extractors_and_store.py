"""Format dispatch, the tabular extractor, and the unit store.

The architectural claim under test: FORMAT (how the file is parsed) and DOMAIN
(which corpus it belongs to) are independent axes. The same tabular extractor
serves financial and medical CSVs, and domain is metadata the caller supplies —
never inferred from the data. If that ever stops holding, adding a third domain
becomes a rewrite instead of a config change.

The narrative generator is stubbed throughout: it costs one Groq call per row,
and a test suite that quietly eats the daily rate limit is worse than none.
"""
import pytest

from src.extractors import tabular_extractor
from src.main import extract_any
from src.unit_store import load_units, merge_units


@pytest.fixture
def stub_llm(monkeypatch):
    """Replace the per-row Groq call with a deterministic string."""
    calls = []

    def fake(row):
        calls.append(row)
        return f"Narrative for {row.get('row_id')}."

    monkeypatch.setattr(tabular_extractor, "generate_narrative", fake)
    return calls


@pytest.fixture
def csv_file(tmp_path):
    path = tmp_path / "samples.csv"
    path.write_text(
        "sample_id,score,notes\nTC-014,0.87,strong\nTC-015,0.12,\nTC-016,0.55,moderate\n",
        encoding="utf-8",
    )
    return str(path)


@pytest.fixture
def store(tmp_path, monkeypatch):
    import src.unit_store as unit_store

    monkeypatch.setattr(unit_store, "GENERATED_DIR", str(tmp_path / "generated"))
    return tmp_path


class TestFormatDispatch:
    def test_csv_routes_to_the_tabular_extractor(self, csv_file, stub_llm):
        units = extract_any(csv_file, "medical")
        assert len(units) == 3
        assert all(u.source_type == "tabular" for u in units)

    def test_docx_routes_to_the_text_extractor(self, tmp_path, monkeypatch):
        import docx

        path = tmp_path / "doc.docx"
        document = docx.Document()
        document.add_paragraph("Body text.")
        document.save(path)

        called = {}
        monkeypatch.setattr(
            "src.main.extract_text", lambda p: called.setdefault("path", p) or []
        )
        extract_any(str(path), "medical")
        assert called["path"] == str(path)

    def test_docx_with_a_non_medical_domain_is_refused(self, tmp_path):
        """Better to fail loudly than to mislabel the corpus."""
        import docx

        path = tmp_path / "doc.docx"
        docx.Document().save(path)
        with pytest.raises(ValueError, match="medical"):
            extract_any(str(path), "financial")

    def test_unsupported_extension_is_refused(self, tmp_path):
        path = tmp_path / "notes.pdf"
        path.write_text("x", encoding="utf-8")
        with pytest.raises(ValueError, match="Unsupported"):
            extract_any(str(path), "medical")


class TestDomainIsAParameter:
    def test_the_same_file_can_be_tagged_either_domain(self, csv_file, stub_llm):
        """The cross-domain proof: identical code path, different tag."""
        med = tabular_extractor.extract_tabular(csv_file, domain="medical")
        fin = tabular_extractor.extract_tabular(csv_file, domain="financial")
        assert [u.text for u in med] == [u.text for u in fin]
        assert {u.domain for u in med} == {"medical"}
        assert {u.domain for u in fin} == {"financial"}

    def test_domain_is_never_defaulted(self, csv_file, stub_llm):
        with pytest.raises(TypeError):
            tabular_extractor.extract_tabular(csv_file)

    def test_invalid_domain_is_refused(self, csv_file, stub_llm):
        with pytest.raises(ValueError, match="domain"):
            tabular_extractor.extract_tabular(csv_file, domain="legal")


class TestSchemaAgnosticism:
    def test_completely_different_column_sets_both_work(self, tmp_path, stub_llm):
        """RULES section 3, non-negotiable 8: no hardcoded columns."""
        stocks = tmp_path / "stocks.csv"
        stocks.write_text("Ticker,Growth\nJGCHEM,0.15\n", encoding="utf-8")
        antigens = tmp_path / "antigens.csv"
        antigens.write_text("sample_id,panel,score\nTC-014,B,0.87\n", encoding="utf-8")

        fin = tabular_extractor.extract_tabular(str(stocks), domain="financial")
        med = tabular_extractor.extract_tabular(str(antigens), domain="medical")
        assert "Ticker" in fin[0].metadata.extra
        assert "sample_id" in med[0].metadata.extra

    def test_row_id_is_excluded_from_extra(self, csv_file, stub_llm):
        """row_id is a loader implementation detail, not source data."""
        units = tabular_extractor.extract_tabular(csv_file, domain="medical")
        assert "row_id" not in units[0].metadata.extra

    def test_ids_are_unique_across_two_files_in_one_domain(self, tmp_path, stub_llm):
        """The exact scenario that used to lose a whole file on upsert."""
        a = tmp_path / "tcell.csv"
        a.write_text("id\n1\n2\n", encoding="utf-8")
        b = tmp_path / "human.csv"
        b.write_text("id\n1\n2\n", encoding="utf-8")

        ids = [u.id for f in (a, b)
               for u in tabular_extractor.extract_tabular(str(f), domain="medical")]
        assert len(ids) == len(set(ids)) == 4


class TestUnitStore:
    def test_merge_accumulates_across_files(self, store, make_unit):
        merge_units([make_unit("a", source_file="a.csv", index=i) for i in range(3)], "financial")
        all_units, added, updated = merge_units(
            [make_unit("b", source_file="b.csv", index=i) for i in range(2)], "financial"
        )
        assert (len(all_units), added, updated) == (5, 2, 0)

    def test_reingest_updates_rather_than_duplicating(self, store, make_unit):
        merge_units([make_unit("original", source_file="a.csv", index=0)], "financial")
        all_units, added, updated = merge_units(
            [make_unit("revised", source_file="a.csv", index=0)], "financial"
        )
        assert (len(all_units), added, updated) == (1, 0, 1)
        assert all_units[0].text == "revised"

    def test_domains_are_stored_separately(self, store, make_unit):
        merge_units([make_unit("f", domain="financial", index=0)], "financial")
        merge_units([make_unit("m", domain="medical", source_file="m.csv", index=0)], "medical")
        assert len(load_units("financial")) == 1
        assert len(load_units("medical")) == 1

    def test_missing_store_returns_empty(self, store):
        assert load_units("financial") == []

    def test_units_survive_the_json_round_trip(self, store, make_unit):
        merge_units([make_unit("text", extra={"ticker": "JGCHEM", "n": 42})], "financial")
        restored = load_units("financial")[0]
        assert restored.text == "text"
        assert restored.metadata.extra == {"ticker": "JGCHEM", "n": 42}
